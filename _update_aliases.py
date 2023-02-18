import os
import re
import time
import json
from collections import defaultdict

import requests
from bs4 import BeautifulSoup

cookie_value = ''


class TengXunDocument:

    def __init__(self, document_url, local_pad_id, cookie_value):
        # excel文档地址
        self.document_url = document_url
        # 此值每一份腾讯文档有一个,需要手动获取
        self.localPadId = local_pad_id
        self.headers = {
            'content-type': 'application/x-www-form-urlencoded',
            'Cookie': cookie_value
        }

    def get_now_user_index(self):
        """
        # 获取当前用户信息,供创建下载任务使用
        :return:
            # nowUserIndex = '4883730fe8b94fbdb94da26a9a63b688'
            # uid = '144115225804776585'
            # utype = 'wx'
        """
        response_body = requests.get(url=self.document_url, headers=self.headers, verify=False)
        parser = BeautifulSoup(response_body.content, 'html.parser')
        global_multi_user_list = re.findall(re.compile('window.global_multi_user=(.*?);'), str(parser))
        if global_multi_user_list:
            user_dict = json.loads(global_multi_user_list[0])
            print(user_dict)
            return user_dict['nowUserIndex']
        return 'cookie过期,请重新输入'

    def export_excel_task(self, export_excel_url):
        """
        导出excel文件任务,供查询文件数据准备进度
        :return:
        """
        body = {
            'docId': self.localPadId, 'version': '2'
        }

        res = requests.post(url=export_excel_url,
                                      headers=self.headers, data=body, verify=False)
        operation_id = res.json()['operationId']
        return operation_id

    def download_excel(self, check_progress_url, file_name):
        """
        下载excel文件
        :return:
        """
        # 拿到下载excel文件的url
        start_time = time.time()
        file_url = ''
        while True:
            res = requests.get(url=check_progress_url, headers=self.headers, verify=False)
            progress = res.json()['progress']
            if progress == 100:
                file_url = res.json()['file_url']
                break
            elif time.time() - start_time > 30:
                print("数据准备超时,请排查")
                break
        if file_url:
            self.headers['content-type'] = 'application/octet-stream'
            res = requests.get(url=file_url, headers=self.headers, verify=False)
            with open(file_name, 'wb') as f:
                f.write(res.content)
            print('下载成功,文件名: ' + file_name)
        else:
            print("下载文件地址获取失败, 下载excel文件不成功")


def download_aliases_file():
    # excel文档地址
    document_url = 'https://docs.qq.com/sheet/DVnJUb0pYeXJxakVk'
    # 此值每一份腾讯文档有一个,需要手动获取
    local_pad_id = '300000000$VrToJXyrqjEd'
    # 打开腾讯文档后,从抓到的接口中获取cookie信息
    tx = TengXunDocument(document_url, local_pad_id, cookie_value)
    now_user_index = tx.get_now_user_index()
    # 导出文件任务url
    export_excel_url = f'https://docs.qq.com/v1/export/export_office?u={now_user_index}'
    # 获取导出任务的操作id
    operation_id = tx.export_excel_task(export_excel_url)
    check_progress_url = f'https://docs.qq.com/v1/export/query_progress?u={now_user_index}&operationId={operation_id}'
    file_name = f'aliases.xlsx'
    tx.download_excel(check_progress_url, file_name)


def music_aliases():
    _music_aliases = defaultdict(list)
    _music_aliases_reverse = defaultdict(list)
    with open(os.path.join('static', 'aliases.csv'), 'r', encoding='utf-8') as f:
        _music_aliases_lines = [l for l in f.readlines() if l.strip()]
    for l in _music_aliases_lines:
        arr = l.strip().split('\t')
        for i in range(len(arr)):
            if arr[i] != '':
                _music_aliases[arr[i].lower()].append(arr[0])
                _music_aliases_reverse[arr[0]].append(arr[i].lower())
    return _music_aliases, _music_aliases_reverse, _music_aliases_lines


def merge_aliases():
    from openpyxl import load_workbook
    try:
        wb = load_workbook('aliases.xlsx')
        ws = wb['舞萌DX']
    except KeyError:
        print('aliases文件错误，请检查文件完整性')
        return

    _, _music_aliases_reverse, _music_aliases_lines = music_aliases()
    _music_titles = list(_music_aliases_reverse.keys())

    for i, rows in enumerate(ws.rows):
        if i == 0: continue
        name = rows[1].value
        aliases = []
        for cell in rows[2:]:
            value = cell.value
            if isinstance(value, (int, float)):
                value = str(int(value))
            if value:
                aliases.append(value.strip())
        if len(aliases) > 1:
            if name in _music_aliases_reverse:
                addition = []
                for a in aliases:
                    if a.lower() not in _music_aliases_reverse[name]:
                        addition.append(a)
                if len(addition) > 1:
                    addition = '\t'.join(addition)
                    _music_aliases_lines[_music_titles.index(name)] = _music_aliases_lines[_music_titles.index(name)].strip() + f'\t{addition}\n'
                    print('add', name, addition)
            else:
                addition = '\t'.join(aliases)
                _music_aliases_lines.append(f'{name}\t{addition}\n')
                print('new', name, addition)
    with open(os.path.join('static', 'aliases.csv'), 'w', encoding='utf-8') as f:
        f.write('\n'.join([l.strip() for l in _music_aliases_lines]))
    print('合并成功')


if __name__ == '__main__':
    download_aliases_file()

    merge_aliases()
